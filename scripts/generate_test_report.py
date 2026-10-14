import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import datetime
import os

TEMPLATE_PATH = r'docs/others/(template) Test Summary Report - HumanResource.xlsx'
OUTPUT_PATH = r'docs/testing/AMIGA_Robot_Admin_Test_Summary_Report.xlsx'

print('Loading template workbook...')
wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

font_regular = Font(name='Arial', size=11, bold=False, color='000000')
font_bold = Font(name='Arial', size=11, bold=True, color='000000')

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

# -------------------------------------------------------------
# 1. UPDATE DASHBOARD SUMMARY
# -------------------------------------------------------------
print('Updating Dashboard Summary...')
ws_dash = wb['Dashboard Summary']

# Title and Logo
ws_dash['A1'].value = 'AMIGA ROBOT ADMIN'
ws_dash['A7'].value = (
    'AMIGA Robot Admin merupakan sistem cerdas pemrosesan invoice vendor (PDF digital maks. 5 halaman) '
    'menjadi voucher pembayaran resmi terverifikasi secara instan. Mengintegrasikan formulir intake terisolasi multi-tenant, '
    'antrean background worker BullMQ, ekstraksi asinkron AWS Bedrock Titan Express, verifikasi 3 aturan bisnis '
    '(Missing Evidence, Possible Duplicate, Calculation Discrepancy), layar peninjauan split-screen dual-context 50:50, '
    'penegakan Segregation of Duties (SoD) mutlak di sisi server, kompilasi voucher PDF resmi (@react-pdf/renderer), '
    'serta pencatatan jejak audit permanen yang tidak dapat diubah (immutable audit ledger).'
)

ws_dash['A13'].value = 'Dokumen Eksekusi Pengujian Sistem & UAT - AMIGA Robot Admin'
ws_dash['A18'].value = 'AMIGA Robot Admin (MAIN Hackathon 2026 PoC)'

ws_dash['A26'].value = 'Tech Lead & QA Engineering Team'
ws_dash['C26'].value = datetime.datetime(2026, 10, 14, 0, 0)
ws_dash['E26'].value = datetime.datetime(2026, 10, 14, 0, 0)
ws_dash['G26'].value = datetime.datetime(2026, 10, 16, 0, 0)

ws_dash['A32'].value = (
    'Memastikan seluruh fitur inti AMIGA Robot Admin berfungsi dengan integritas finansial tinggi dan bebas cacat: '
    '1) Intake & validasi format berkas PDF digital (maks. 5 halaman, 10MB) serta penerbitan Request ID berstatus QUEUED; '
    '2) Ekstraksi teks in-context AI AWS Bedrock Titan Express ke struktur data voucher; '
    '3) Mesin verifikasi otomatis 3 aturan sistem (Missing Evidence, Possible Duplicate, Calculation Discrepancy); '
    '4) Dashboard antrean kerja multi-klien dan pemantauan batas beban reviewer (>10); '
    '5) Layar peninjauan split-screen 50:50 dengan kontrol PDF canvas dan editor finansial kalkulasi real-time; '
    '6) Penegakan Segregation of Duties (SoD) di server dan auto-reassignment ke least-loaded active reviewer; '
    '7) Kompilasi dan unduh voucher PDF resmi (SLA <= 3s) dengan teks otorisasi permanen; '
    '8) Batas isolasi ketat multi-tenant dan keabadian jejak audit (immutable audit trail).'
)

ws_dash['G32'].value = (
    'Dev / Local Environment:\n'
    '• Frontend: http://localhost:3000 (Next.js 16 LTS, React 19, Tailwind v4)\n'
    '• Backend: http://localhost:4000 (Fastify, Prisma ORM, PostgreSQL 16)\n'
    '• Message Broker: Redis 7.2 (BullMQ Worker)\n'
    '• AI Engine: AWS Bedrock Titan Text Express (ap-southeast-1) + Deterministic Mock'
)

# Update Defect Contribution Feature Names
features_list = [
    'Authentication & RBAC',
    'Intake & Upload Modul',
    'AI Document Extraction',
    'Verification Engine',
    'Dashboard Antrean',
    'Split-Screen Review',
    'Segregation of Duties',
    'Voucher Generator',
    'Document Correction',
    'Multi-Tenant & Audit'
]

for idx, feat in enumerate(features_list):
    row = 54 + idx
    ws_dash[f'J{row}'].value = feat
    ws_dash[f'L{row}'].value = f'=COUNTIF(\'Defect Tracking Data\'!E9:E500, J{row})'

ws_dash['M54'].value = f'=SUM(L54:L{54 + len(features_list) - 1})'

# Update formulas in Dashboard Summary to ensure robust ranges
ws_dash['G45'].value = '=COUNTIF(\'System Testing - Execution Data\'!Q9:R500,"Passed") + COUNTIF(\'End-to-End Testing - Execution \'!Q9:R500, "Passed")'
ws_dash['G47'].value = '=COUNTIF(\'System Testing - Execution Data\'!Q9:R500,"Failed") + COUNTIF(\'End-to-End Testing - Execution \'!Q9:R500, "Failed")'
ws_dash['G49'].value = '=COUNTIF(\'System Testing - Execution Data\'!Q9:R500,"Blocked") + COUNTIF(\'End-to-End Testing - Execution \'!Q9:R500, "Blocked")'
ws_dash['G51'].value = '=COUNTIF(\'System Testing - Execution Data\'!Q9:R500,"Not Executed") + COUNTIF(\'End-to-End Testing - Execution \'!Q9:R500, "Not Executed")'
ws_dash['G43'].value = '=SUM(G45:G52)'
ws_dash['E18'].value = '=G43'
ws_dash['E22'].value = '=IFERROR(G45/G43, 0)'

ws_dash['C45'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"New")'
ws_dash['C47'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Open")'
ws_dash['C49'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Fixed")'
ws_dash['C51'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Pending Retest")'
ws_dash['C53'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Retest")'
ws_dash['C55'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Reopen")'
ws_dash['C57'].value = '=COUNTIF(\'Defect Tracking Data\'!M9:M500,"Closed")'

ws_dash['C59'].value = '=COUNTIF(\'Defect Tracking Data\'!O9:O500,"Minor")'
ws_dash['C61'].value = '=COUNTIF(\'Defect Tracking Data\'!O9:O500,"Major")'
ws_dash['C63'].value = '=COUNTIF(\'Defect Tracking Data\'!O9:O500,"Critical")'
ws_dash['C43'].value = '=SUM(C45:C57)'
ws_dash['E13'].value = '=C43'


# -------------------------------------------------------------
# HELPER: CLEAR SHEET ROWS & POPULATE TEST CASES
# -------------------------------------------------------------
def reset_sheet_data(ws, start_row=9):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)


def write_test_row(ws, row_idx, test_item):
    """
    Writes a formatted test case row matching the template layout:
    A:B = Main Feature
    C:D = Sub-Feature
    E = Test ID
    F = Date Tested
    G:H = Scenario Name / Test Case
    I:J = Description
    K:L = Precondition
    M:N = Test Steps
    O:P = Expected Result
    Q:R = Execution Status
    S:T = Actual Result
    U:V = Priority
    W:X = Test Data
    Y:Z = Screenshot/Link Source
    AA:AB = Tested By
    AC:AD = Remarks/Comments
    """
    r = row_idx
    
    # Values mapping to the first cell of each merged pair
    cells_data = {
        'A': test_item.get('main_feature', ''),
        'C': test_item.get('sub_feature', ''),
        'E': test_item.get('test_id', ''),
        'F': test_item.get('date_tested', datetime.datetime(2026, 10, 14, 0, 0)),
        'G': test_item.get('scenario_name', ''),
        'I': test_item.get('description', ''),
        'K': test_item.get('precondition', ''),
        'M': test_item.get('test_steps', ''),
        'O': test_item.get('expected_result', ''),
        'Q': test_item.get('execution_status', 'Not Executed'),
        'S': test_item.get('actual_result', 'Siap dieksekusi'),
        'U': test_item.get('priority', 'High'),
        'W': test_item.get('test_data', '-'),
        'Y': test_item.get('screenshot_link', '-'),
        'AA': test_item.get('tested_by', 'QA Engineer'),
        'AC': test_item.get('remarks', '-')
    }
    
    # Merge column pairs for this row
    pairs = [
        ('A', 'B'), ('C', 'D'), ('G', 'H'), ('I', 'J'),
        ('K', 'L'), ('M', 'N'), ('O', 'P'), ('Q', 'R'),
        ('S', 'T'), ('U', 'V'), ('W', 'X'), ('Y', 'Z'),
        ('AA', 'AB'), ('AC', 'AD')
    ]
    for c1, c2 in pairs:
        ws.merge_cells(f'{c1}{r}:{c2}{r}')
        
    for col_letter in ['A', 'C', 'E', 'F', 'G', 'I', 'K', 'M', 'O', 'Q', 'S', 'U', 'W', 'Y', 'AA', 'AC']:
        ws[f'{col_letter}{r}'].value = cells_data[col_letter]
        
    # Apply styling and borders to all 30 columns in this row
    center_cols = {'A', 'B', 'C', 'D', 'E', 'Q', 'R', 'U', 'V'}
    for col_idx in range(1, 31):
        c = ws.cell(r, col_idx)
        col_letter = get_column_letter(col_idx)
        c.font = font_regular
        c.border = thin_border
        if col_letter in center_cols:
            c.alignment = align_center
        else:
            c.alignment = align_left
            
    # Set status badge fill if Passed / Failed / Not Executed
    status = test_item.get('execution_status', 'Not Executed')
    status_fill = None
    if status == 'Passed':
        status_fill = PatternFill(start_color='D9EAD3', end_color='D9EAD3', fill_type='solid')
    elif status == 'Failed':
        status_fill = PatternFill(start_color='F4CCCC', end_color='F4CCCC', fill_type='solid')
    elif status == 'Blocked':
        status_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    elif status == 'Not Executed':
        status_fill = PatternFill(start_color='EFEFEF', end_color='EFEFEF', fill_type='solid')
        
    if status_fill:
        ws.cell(r, 17).fill = status_fill
        ws.cell(r, 18).fill = status_fill

    ws.row_dimensions[r].height = 45


# -------------------------------------------------------------
# 2. SYSTEM TESTING - EXECUTION DATA TEST CASES (TS-001 to TS-036)
# -------------------------------------------------------------
print('Generating System Testing test cases...')
ws_sys = wb['System Testing - Execution Data']
reset_sheet_data(ws_sys, start_row=9)

system_tests = [
    # --- Authentication & RBAC ---
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'Login PIC Klien',
        'test_id': 'TS-001',
        'scenario_name': 'Login akun PIC Klien dan verifikasi akses modul intake',
        'description': 'Memastikan pengguna PIC Klien berhasil login dan hanya mengakses dashboard antrean, form upload, dan voucher kliennya.',
        'precondition': 'Akun PIC Klien telah terdaftar di database (seeder)',
        'test_steps': '1. Buka http://localhost:3000\n2. Masukkan email pic_client@amiga.co.id & password\n3. Klik tombol Masuk\n4. Amati navigasi navbar dan akses menu',
        'expected_result': 'Login berhasil, token JWT terbit, diarahkan ke Dashboard antrean dengan tombol Unggah Invoice aktif.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Email: pic_client@amiga.co.id\nPassword: password123!',
        'remarks': 'Memenuhi PR-004, PF-001'
    },
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'Login Finance Reviewer',
        'test_id': 'TS-002',
        'scenario_name': 'Login akun Finance Reviewer dan verifikasi akses review',
        'description': 'Memastikan Finance Reviewer dapat login dan mengakses antrean dokumen serta aksi Review pada layar split-screen.',
        'precondition': 'Akun Reviewer telah terdaftar di database',
        'test_steps': '1. Buka halaman login\n2. Masukkan email reviewer1@amiga.co.id & password\n3. Klik Masuk\n4. Verifikasi tombol Review pada baris dokumen READY_FOR_REVIEW',
        'expected_result': 'Reviewer berhasil login, tombol Tindak Lanjut / Review aktif pada tabel antrean.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Email: reviewer1@amiga.co.id\nPassword: password123!',
        'remarks': 'Memenuhi PR-004, PR-015'
    },
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'Login Service Lead',
        'test_id': 'TS-003',
        'scenario_name': 'Login Service Lead dan verifikasi widget monitoring beban',
        'description': 'Memastikan Service Lead dapat melihat antrean lintas reviewer dan bar beban reviewer.',
        'precondition': 'Akun Service Lead telah terdaftar di database',
        'test_steps': '1. Login dengan akun service_lead@amiga.co.id\n2. Masuk ke halaman Queue\n3. Amati Reviewer Workload Bar',
        'expected_result': 'Seluruh antrean reviewer tampil lengkap dengan metrik beban aktif per reviewer.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Email: service_lead@amiga.co.id\nPassword: password123!',
        'remarks': 'Memenuhi PR-023, SR-014'
    },
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'Invalid Credentials',
        'test_id': 'TS-004',
        'scenario_name': 'Login dengan password atau email salah',
        'description': 'Memastikan sistem menolak kredensial salah dengan respon error RFC 7807 401 Unauthorized.',
        'precondition': 'Halaman login terbuka',
        'test_steps': '1. Masukkan email reviewer1@amiga.co.id\n2. Masukkan password salah: passwordsalah999\n3. Klik Masuk',
        'expected_result': 'Sistem menampilkan notifikasi peringatan kredensial salah dan token JWT tidak diterbitkan.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Password: passwordsalah999',
        'remarks': 'Memenuhi SR-012'
    },
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'RBAC Route Guard',
        'test_id': 'TS-005',
        'scenario_name': 'Akses URL terproteksi reviewer oleh akun PIC Klien',
        'description': 'Memastikan akses langsung ke halaman /review/:id oleh akun PIC Klien ditolak sistem.',
        'precondition': 'Login sebagai PIC Klien',
        'test_steps': '1. Buka URL langsung http://localhost:3000/review/REQ-TEST-001\n2. Amati perilaku browser',
        'expected_result': 'Sistem menolak akses dan me-redirect ke Dashboard atau menampilkan status 403 Forbidden.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'URL: /review/[id]',
        'remarks': 'Memenuhi SR-002, SR-003'
    },
    {
        'main_feature': 'Authentication & RBAC',
        'sub_feature': 'System Admin Restriction',
        'test_id': 'TS-006',
        'scenario_name': 'Verifikasi larangan persetujuan transaksi oleh System Admin',
        'description': 'Memastikan peran System Admin dilarang menyetujui transaksi keuangan di tingkat server.',
        'precondition': 'Login sebagai sysadmin@amiga.co.id',
        'test_steps': '1. Coba panggil endpoint POST /api/v1/invoices/:id/approve dengan bearer token sysadmin',
        'expected_result': 'Server menolak request dengan status HTTP 403 Forbidden (Non-negotiable SoD rule).',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Email: sysadmin@amiga.co.id',
        'remarks': 'Memenuhi BR-014, PR-024'
    },

    # --- Intake & Upload Modul ---
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'Valid PDF Upload',
        'test_id': 'TS-007',
        'scenario_name': 'Unggah berkas PDF digital valid (1-5 halaman, <= 10MB)',
        'description': 'Memastikan berkas PDF valid berhasil diunggah dan sistem menerbitkan Request ID unik berstatus QUEUED.',
        'precondition': 'Login sebagai PIC Klien atau Intake Support',
        'test_steps': '1. Buka menu Unggah Invoice (/upload)\n2. Pilih Klien: CLI-MAIN\n3. Pilih Saluran: EMAIL\n4. Drag & drop berkas PDF digital valid\n5. Klik Unggah Dokumen',
        'expected_result': 'Muncul modal sukses memuat Request ID unik (REQ-YYYYMMDD-XXXXX) dan status QUEUED dalam tempo <= 500ms.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'File: invoice_sample_valid.pdf (1 halaman, 120KB)',
        'remarks': 'Memenuhi PR-001, PR-003, SR-008'
    },
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'Magic Bytes Non-PDF Rejection',
        'test_id': 'TS-008',
        'scenario_name': 'Penolakan berkas non-PDF atau ekstensi palsu (.exe diubah .pdf)',
        'description': 'Memastikan server memvalidasi magic bytes header (%PDF-1.) dan menolak berkas manipulasi.',
        'precondition': 'Halaman formulir unggah terbuka',
        'test_steps': '1. Ambil berkas teks / binary (.exe) lalu ubah namanya menjadi invoice.pdf\n2. Unggah berkas tersebut pada formulir\n3. Klik submit',
        'expected_result': 'Server menolak berkas dengan pesan INVALID_FILE_TYPE dan berkas tidak disimpan ke storage.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'File: fake_invoice.pdf (binary non-PDF)',
        'remarks': 'Memenuhi T-005, SR-008'
    },
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'Page Limit Exceeded (>5 Halaman)',
        'test_id': 'TS-009',
        'scenario_name': 'Penolakan berkas PDF melebihi batas 5 halaman',
        'description': 'Memastikan berkas PDF dengan 6 halaman atau lebih ditolak secara otomatis.',
        'precondition': 'Berkas PDF 6 halaman tersedia',
        'test_steps': '1. Unggah berkas PDF 6 halaman\n2. Submit formulir intake',
        'expected_result': 'Sistem menolak dengan kode error PAGE_LIMIT_EXCEEDED dan meminta PIC mengunggah invoice ringkas.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'File: invoice_6_pages.pdf',
        'remarks': 'Memenuhi PR-001, PR-005'
    },
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'File Size Limit (>10MB)',
        'test_id': 'TS-010',
        'scenario_name': 'Penolakan berkas PDF melebihi ukuran 10MB',
        'description': 'Memastikan sistem client dan server membatasi ukuran berkas maksimal 10MB.',
        'precondition': 'Berkas PDF berukuran 12MB tersedia',
        'test_steps': '1. Pilih berkas PDF 12MB pada drag-and-drop area',
        'expected_result': 'Validasi client menampilkan peringatan ukuran maksimal 10MB dan tombol unggah dinonaktifkan.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'File size: 12.4 MB',
        'remarks': 'Memenuhi PR-001'
    },
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'Metadata Completeness Validation',
        'test_id': 'TS-011',
        'scenario_name': 'Validasi kelengkapan metadata formulir intake',
        'description': 'Memastikan formulir mewajibkan pengisian Klien, Saluran, dan Tanggal Pengajuan.',
        'precondition': 'Formulir unggah terbuka',
        'test_steps': '1. Kosongkan pilihan klien atau saluran pengajuan\n2. Klik tombol submit',
        'expected_result': 'Muncul pesan validasi wajib isi pada field yang belum dilengkapi.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'Client ID: Kosong',
        'remarks': 'Memenuhi PR-002'
    },
    {
        'main_feature': 'Intake & Upload Modul',
        'sub_feature': 'Intake Support Delegation Guard',
        'test_id': 'TS-0012',
        'scenario_name': 'Delegasi pengunggahan oleh Intake Support pada klien yang ditugaskan',
        'description': 'Memastikan Intake Support hanya dapat mengunggah untuk klien yang ada dalam daftar tugas resminya.',
        'precondition': 'Login sebagai intake_support@amiga.co.id',
        'test_steps': '1. Buka dropdown Klien pada formulir unggah\n2. Periksa daftar klien yang dapat dipilih',
        'expected_result': 'Hanya organisasi klien yang ditugaskan (assigned clients) yang muncul pada opsi pilihan.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Akun: intake_support@amiga.co.id',
        'remarks': 'Memenuhi PR-004, BR-006'
    },

    # --- AI Document Extraction ---
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'BullMQ Worker Status Transition',
        'test_id': 'TS-013',
        'scenario_name': 'Mutasi status dokumen otomatis ke PROCESSING saat AI membaca berkas',
        'description': 'Memastikan status dokumen berubah dari QUEUED ke PROCESSING saat worker BullMQ mengambil job.',
        'precondition': 'Dokumen berhasil diunggah (QUEUED)',
        'test_steps': '1. Amati status dokumen pada tabel dashboard segera setelah diunggah\n2. Verifikasi state transition di database',
        'expected_result': 'Status dokumen bertransisi menjadi PROCESSING secara otomatis dalam antrean BullMQ.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Queue: ai_extraction_queue',
        'remarks': 'Memenuhi PR-006, PF-007'
    },
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'Header & Vendor Extraction',
        'test_id': 'TS-014',
        'scenario_name': 'Ekstraksi nama vendor, nomor invoice, tanggal dan jatuh tempo oleh AI Bedrock',
        'description': 'Memastikan AI mengekstrak data header vendor secara akurat ke tabel invoice_requests.',
        'precondition': 'Dokumen diproses oleh AI extraction worker',
        'test_steps': '1. Tunggu worker menyelesaikan ekstraksi\n2. Periksa data hasil ekstraksi di database / review page',
        'expected_result': 'Nama vendor, nomor invoice, tanggal invoice, dan jatuh tempo terekstraksi sesuai teks PDF.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Vendor: PT Media Solusi Pratama',
        'remarks': 'Memenuhi BR-007, SR-009'
    },
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'Multi-line Items Extraction',
        'test_id': 'TS-015',
        'scenario_name': 'Ekstraksi 100% baris item pada invoice dengan banyak rincian barang',
        'description': 'Memastikan seluruh baris item (deskripsi, kuantitas, harga satuan, subtotal) terbaca lengkap.',
        'precondition': 'Invoice memiliki 3 atau lebih line items',
        'test_steps': '1. Unggah invoice dengan 3 baris item\n2. Periksa tabel invoice_line_items setelah ekstraksi',
        'expected_result': 'Tepat 3 baris item tersimpan lengkap dengan perhitungan subtotal masing-masing baris.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Items: Cloud Server, Maintenance, Managed Service',
        'remarks': 'Memenuhi BR-004, UAT-2'
    },
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'Payment Bank Account Extraction',
        'test_id': 'TS-016',
        'scenario_name': 'Ekstraksi rekening bank transfer tujuan pembayaran',
        'description': 'Memastikan nomor rekening, nama bank, dan nama pemegang rekening terekstraksi ke invoice_payment_details.',
        'precondition': 'Invoice mencantumkan rekening bank',
        'test_steps': '1. Unggah invoice memuat nomor rekening BCA / Mandiri\n2. Periksa tabel invoice_payment_details',
        'expected_result': 'Nomor rekening dan nama bank terekstraksi secara tepat.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Bank: Bank Central Asia (BCA), No: 5410982711',
        'remarks': 'Memenuhi BR-007'
    },
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'Extraction SLA Performance',
        'test_id': 'TS-017',
        'scenario_name': 'Pengukuran durasi pemrosesan ekstraksi AI (SLA <= 15 detik)',
        'description': 'Memastikan proses ekstraksi dokumen hingga 5 halaman selesai dalam tempo tidak melebihi 15 detik.',
        'precondition': 'Job masuk ke worker',
        'test_steps': '1. Catat timestamp saat status QUEUED\n2. Catat timestamp saat hasil ekstraksi tersimpan',
        'expected_result': 'Total durasi ekstraksi <= 15.000 ms tercatat pada log backend terstruktur.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'SLA target: <= 15 detik',
        'remarks': 'Memenuhi BR-003, PR-008'
    },
    {
        'main_feature': 'AI Document Extraction',
        'sub_feature': 'Mock Fallback on Bedrock Failure',
        'test_id': 'TS-018',
        'scenario_name': 'Aktivasi fallback deterministik saat koneksi Bedrock terputus / timeout',
        'description': 'Memastikan dokumen tidak gagal total/hilang jika API Bedrock offline; fallback mock mempertahankan data.',
        'precondition': 'Koneksi Bedrock disimulasikan error / mock mode',
        'test_steps': '1. Jalankan worker dengan fallback enabled\n2. Unggah invoice uji',
        'expected_result': 'Sistem tetap mengekstrak data uji secara deterministik tanpa crash unhandled exception.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'MOCK_BEDROCK=true',
        'remarks': 'Memenuhi PR-010, SR-009'
    },

    # --- Verification Engine ---
    {
        'main_feature': 'Verification Engine',
        'sub_feature': 'Clean Invoice Verification',
        'test_id': 'TS-019',
        'scenario_name': 'Invoice bersih tanpa kendala otomatis bertransisi ke READY_FOR_REVIEW',
        'description': 'Memastikan dokumen dengan rekening lengkap, bukan duplikat, dan kalkulasi tepat lolos verifikasi.',
        'precondition': 'Ekstraksi AI selesai untuk invoice valid',
        'test_steps': '1. Eksekusi verification service pada invoice bersih\n2. Cek status akhir di database',
        'expected_result': 'Status dokumen bertransisi menjadi READY_FOR_REVIEW dan tidak ada catatan di verification_issues.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Invoice bersih lengkap',
        'remarks': 'Memenuhi PR-012, PF-009'
    },
    {
        'main_feature': 'Verification Engine',
        'sub_feature': 'Missing Evidence Detection',
        'test_id': 'TS-020',
        'scenario_name': 'Deteksi invoice tanpa nomor rekening bank (MISSING_EVIDENCE)',
        'description': 'Memastikan dokumen tanpa rekening bank otomatis ditandai kendala dan bertransisi ke status EXCEPTION.',
        'precondition': 'Unggah berkas invoice yang sengaja tidak mencantumkan data bank transfer',
        'test_steps': '1. Tunggu proses ekstraksi dan verifikasi selesai\n2. Periksa status invoice dan tabel verification_issues',
        'expected_result': 'Status invoice menjadi EXCEPTION dengan issueCode MISSING_EVIDENCE dan deskripsi informatif.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Rekening bank: Kosong',
        'remarks': 'Memenuhi BR-008, UAT-3'
    },
    {
        'main_feature': 'Verification Engine',
        'sub_feature': 'Duplicate Invoice Detection',
        'test_id': 'TS-021',
        'scenario_name': 'Deteksi invoice ganda pada klien yang sama (POSSIBLE_DUPLICATE)',
        'description': 'Memastikan sistem mendeteksi invoice dengan nomor dan vendor yang identik setelah normalisasi teks.',
        'precondition': 'Invoice INV-2026-001 dari Vendor A telah tersimpan sebelumnya pada klien CLI-MAIN',
        'test_steps': '1. Unggah kembali invoice dengan vendor dan nomor INV-2026-001 pada klien CLI-MAIN\n2. Amati hasil verifikasi',
        'expected_result': 'Sistem menandai isu POSSIBLE_DUPLICATE dan memutasi status menjadi EXCEPTION.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Invoice No: INV-2026-001 (Duplicate)',
        'remarks': 'Memenuhi BR-008, SR-013, UAT-4'
    },
    {
        'main_feature': 'Verification Engine',
        'sub_feature': 'Calculation Discrepancy Detection',
        'test_id': 'TS-022',
        'scenario_name': 'Deteksi selisih subtotal + pajak != grand total (CALCULATION_DISCREPANCY)',
        'description': 'Memastikan sistem mendeteksi ketidaksesuaian matematis pada baris item terhadap grand total.',
        'precondition': 'Invoice memiliki total Rp 1.500.000 namun jumlah baris item hanya Rp 1.300.000',
        'test_steps': '1. Proses invoice dengan selisih nominal\n2. Amati hasil verifikasi sistem',
        'expected_result': 'Status menjadi EXCEPTION dengan issueCode CALCULATION_DISCREPANCY.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Discrepancy: Rp 200.000',
        'remarks': 'Memenuhi BR-008, PR-011'
    },

    # --- Document Exception Correction ---
    {
        'main_feature': 'Document Correction',
        'sub_feature': 'Correction by PIC Client',
        'test_id': 'TS-023',
        'scenario_name': 'Koreksi data rekening dokumen EXCEPTION oleh PIC dengan Request ID identik',
        'description': 'Memastikan PIC dapat mengirimkan perbaikan rekening tanpa mengubah Request ID dan memicu verifikasi ulang.',
        'precondition': 'Dokumen berstatus EXCEPTION karena MISSING_EVIDENCE',
        'test_steps': '1. Login sebagai PIC Klien\n2. Buka modal koreksi dokumen\n3. Masukkan data rekening bank yang benar\n4. Klik Simpan Perbaikan',
        'expected_result': 'Request ID TETAP IDENTIK, status berubah ke PROCESSING lalu lolos menjadi READY_FOR_REVIEW.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Bank: BCA, No: 1234567890',
        'remarks': 'Memenuhi PR-009, PR-014, UAT-8'
    },
    {
        'main_feature': 'Document Correction',
        'sub_feature': 'Invalid Correction Prevention',
        'test_id': 'TS-024',
        'scenario_name': 'Penolakan koreksi pada dokumen yang bukan berstatus EXCEPTION',
        'description': 'Memastikan endpoint koreksi menolak request jika status dokumen saat ini READY_FOR_REVIEW atau AMIGA_REVIEWED.',
        'precondition': 'Dokumen berstatus READY_FOR_REVIEW',
        'test_steps': '1. Coba panggil PUT /api/v1/invoices/:id/correct pada dokumen READY_FOR_REVIEW',
        'expected_result': 'Server menolak request dengan status HTTP 400 Bad Request.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'Status target: READY_FOR_REVIEW',
        'remarks': 'Memenuhi T-011'
    },

    # --- Dashboard Antrean & Workload Monitoring ---
    {
        'main_feature': 'Dashboard Antrean',
        'sub_feature': 'Queue Filter & Tabs',
        'test_id': 'TS-025',
        'scenario_name': 'Penyaringan dokumen antrean berdasarkan status dan klien',
        'description': 'Memastikan tab status (Semua, Queued, Processing, Ready, Exception, Approved) dan dropdown filter bekerja presisi.',
        'precondition': 'Terdapat data invoice di berbagai status',
        'test_steps': '1. Buka halaman Queue\n2. Klik tab EXCEPTION\n3. Pilih filter Klien CLI-MAIN',
        'expected_result': 'Hanya dokumen berstatus EXCEPTION milik klien CLI-MAIN yang tampil pada tabel.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'Filter: EXCEPTION & CLI-MAIN',
        'remarks': 'Memenuhi PR-023, DR-001'
    },
    {
        'main_feature': 'Dashboard Antrean',
        'sub_feature': 'Queue Load Time SLA',
        'test_id': 'TS-026',
        'scenario_name': 'Waktu render antrean dashboard (SLA <= 2 detik)',
        'description': 'Memastikan halaman dashboard termuat lengkap dengan auto-refetch TanStack Query dalam tempo <= 2s.',
        'precondition': 'Database memuat minimal 20 records dokumen',
        'test_steps': '1. Hard refresh browser pada halaman /queue\n2. Ukur waktu hingga tabel dan metric cards tampil lengkap',
        'expected_result': 'Halaman antrean selesai dirender dalam durasi <= 2.0 detik.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'SLA Target: <= 2 detik',
        'remarks': 'Memenuhi BR-003, SR-005'
    },
    {
        'main_feature': 'Dashboard Antrean',
        'sub_feature': 'Reviewer Workload Threshold Warning',
        'test_id': 'TS-027',
        'scenario_name': 'Peringatan visual batas beban kerja reviewer (> WORKLOAD_MAX_THRESHOLD)',
        'description': 'Memastikan badge merah peringatan muncul jika antrean aktif reviewer melebihi ambang batas 10 dokumen.',
        'precondition': 'Reviewer 1 ditugaskan 11 dokumen aktif',
        'test_steps': '1. Login sebagai Service Lead\n2. Amati widget Reviewer Workload Bar',
        'expected_result': 'Bar Reviewer 1 menampilkan indikator peringatan overload warna merah.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Threshold: > 10 dokumen',
        'remarks': 'Memenuhi PR-023, SR-014'
    },
    {
        'main_feature': 'Dashboard Antrean',
        'sub_feature': 'Service Lead Manual Reassignment',
        'test_id': 'TS-028',
        'scenario_name': 'Pengalihan penugasan dokumen oleh Service Lead',
        'description': 'Memastikan Service Lead dapat memindahkan penugasan dokumen antar-reviewer.',
        'precondition': 'Login sebagai Service Lead',
        'test_steps': '1. Buka dropdown aksi penugasan pada dokumen di antrean\n2. Pilih Alihkan Penugasan ke Reviewer 2\n3. Konfirmasi modal pengalihan',
        'expected_result': 'Penugasan dokumen berpindah ke Reviewer 2 dan tercatat pada tabel reviewer_assignments.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Target: reviewer2@amiga.co.id',
        'remarks': 'Memenuhi PR-023, T-017'
    },

    # --- Split-Screen Review & Financial Editor ---
    {
        'main_feature': 'Split-Screen Review',
        'sub_feature': 'Dual-Context Split-Screen 50:50',
        'test_id': 'TS-029',
        'scenario_name': 'Tampilan antarmuka split-screen 50:50 PDF canvas dan form editor',
        'description': 'Memastikan layar review menampilkan berkas PDF asli di panel kiri dan form data di panel kanan tanpa layout shift.',
        'precondition': 'Login sebagai Finance Reviewer dan buka dokumen READY_FOR_REVIEW',
        'test_steps': '1. Klik tombol Review pada dokumen\n2. Amati struktur tampilan layar',
        'expected_result': 'Layar terbagi 50:50 dengan PDF.js viewer di kiri dan financial editor di kanan, waktu muat <= 2 detik.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Layout: CSS Grid 50:50',
        'remarks': 'Memenuhi PR-015, DR-003'
    },
    {
        'main_feature': 'Split-Screen Review',
        'sub_feature': 'PDF Viewer Controls',
        'test_id': 'TS-030',
        'scenario_name': 'Kontrol zoom, fit-width, dan navigasi halaman PDF canvas',
        'description': 'Memastikan fungsi zoom (50%-200%), fit-width, dan navigasi halaman 1 s/d 5 berjalan mulus dan bebas blur.',
        'precondition': 'Layar split-screen terbuka untuk dokumen multi-halaman',
        'test_steps': '1. Klik tombol Zoom In (+)\n2. Klik tombol Zoom Out (-)\n3. Pindah ke halaman 2\n4. Klik tombol Fit Width',
        'expected_result': 'Canvas PDF menyesuaikan zoom dan halaman berganti secara responsif tanpa flickering.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'Zoom range: 50% - 200%',
        'remarks': 'Memenuhi T-014'
    },
    {
        'main_feature': 'Split-Screen Review',
        'sub_feature': 'Interactive Subtotal Recalculation',
        'test_id': 'TS-031',
        'scenario_name': 'Kalkulasi ulang matematis real-time saat baris item diedit',
        'description': 'Memastikan perubahan harga satuan atau kuantitas langsung memperbarui subtotal baris dan grand total.',
        'precondition': 'Form line items editor terbuka di panel kanan',
        'test_steps': '1. Ubah harga satuan baris 1 dari Rp 10.000 menjadi Rp 15.000 (kuantitas 2)\n2. Tekan Tab / klik di luar input',
        'expected_result': 'Subtotal baris 1 seketika berubah menjadi Rp 30.000 dan grand total terkalkulasi otomatis.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Qty: 2, Unit Price: 15.000',
        'remarks': 'Memenuhi PR-016, PF-014'
    },
    {
        'main_feature': 'Split-Screen Review',
        'sub_feature': 'Return Document to PIC',
        'test_id': 'TS-032',
        'scenario_name': 'Pengembalian dokumen ke PIC Klien dengan catatan justifikasi wajib',
        'description': 'Memastikan reviewer dapat mengembalikan dokumen bermasalah ke status EXCEPTION dengan alasan tertulis.',
        'precondition': 'Reviewer membuka dokumen READY_FOR_REVIEW',
        'test_steps': '1. Klik tombol Kembalikan ke PIC\n2. Tulis alasan: Fisik invoice buram dan stempel menutupi rekening\n3. Klik Konfirmasi Pengembalian',
        'expected_result': 'Status dokumen berubah menjadi EXCEPTION dan catatan pengembalian tersimpan di audit log.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Alasan pengembalian wajib diisi',
        'remarks': 'Memenuhi PR-017'
    },

    # --- Segregation of Duties (SoD) & Auto-Reassignment ---
    {
        'main_feature': 'Segregation of Duties',
        'sub_feature': 'Atomic Server-Side SoD Guard',
        'test_id': 'TS-033',
        'scenario_name': 'Penolakan persetujuan dokumen oleh pengunggah di tingkat database atomik',
        'description': 'Memastikan jika Reviewer == Uploader menekan tombol Setujui, transaksi server dibatalkan dan memicu SoD violation.',
        'precondition': 'Reviewer 1 mengunggah dokumen lalu mencoba menyetujui dokumen tersebut',
        'test_steps': '1. Login sebagai reviewer1@amiga.co.id\n2. Unggah dokumen invoice\n3. Buka halaman review dokumen tersebut\n4. Klik tombol Setujui Dokumen',
        'expected_result': 'Transaksi ditolak server HTTP 403 SOD_VIOLATION, persetujuan dibatalkan, status tidak menjadi AMIGA_REVIEWED.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'uploader_id == reviewer_id',
        'remarks': 'Memenuhi BR-002, SR-004, UAT-5'
    },
    {
        'main_feature': 'Segregation of Duties',
        'sub_feature': 'Auto-Reassignment to Least-Loaded Reviewer',
        'test_id': 'TS-034',
        'scenario_name': 'Pengalihan otomatis ke reviewer paling longgar saat SoD dilanggar',
        'description': 'Memastikan sistem otomatis mengalihkan dokumen ke reviewer aktif dengan antrean tersedikit saat SoD terjadi.',
        'precondition': 'Reviewer 2 aktif dan memiliki antrean lebih sedikit dibanding reviewer lainnya',
        'test_steps': '1. Picu SoD violation pada dokumen\n2. Periksa field current_reviewer_id dan tabel reviewer_assignments',
        'expected_result': 'Dokumen dialihkan ke Reviewer 2, status tetap READY_FOR_REVIEW, dan event SOD_AUTO_REASSIGN tercatat di audit.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Least-loaded reviewer logic',
        'remarks': 'Memenuhi BR-002, PR-018, T-017'
    },

    # --- Voucher Generation & Download ---
    {
        'main_feature': 'Voucher Generator',
        'sub_feature': 'Voucher PDF Compilation SLA',
        'test_id': 'TS-035',
        'scenario_name': 'Kompilasi voucher PDF resmi pasca persetujuan reviewer (SLA <= 3 detik)',
        'description': 'Memastikan voucher resmi dikompilasi (@react-pdf/renderer) dalam tempo <= 3s dengan nomor VCH/YYYYMM/CLI/XXXXX.',
        'precondition': 'Dokumen disetujui oleh reviewer independen yang valid',
        'test_steps': '1. Klik Setujui Dokumen sebagai reviewer independen\n2. Amati waktu pembentukan voucher dan status dokumen',
        'expected_result': 'Status menjadi AMIGA_REVIEWED, nomor voucher terbit unik, dan berkas PDF voucher tersimpan di storage dalam durasi <= 3s.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'SLA Target: <= 3 detik',
        'remarks': 'Memenuhi BR-003, PR-020, SR-010'
    },
    {
        'main_feature': 'Voucher Generator',
        'sub_feature': 'Legal Authorization Text & Download',
        'test_id': 'TS-036',
        'scenario_name': 'Verifikasi teks batas otorisasi legal dan mutasi status ke SENT_TO_CLIENT saat unduh',
        'description': 'Memastikan voucher memuat klausul batas legalitas dan pengunduhan oleh PIC mengubah status ke SENT_TO_CLIENT.',
        'precondition': 'Dokumen berstatus AMIGA_REVIEWED',
        'test_steps': '1. Login sebagai PIC Klien\n2. Buka layar pratinjau voucher (/voucher/:id)\n3. Verifikasi teks batas otorisasi legal tertera verbatim\n4. Klik Unduh Voucher PDF',
        'expected_result': 'Browser mengunduh berkas PDF voucher dan status dokumen di database otomatis bermutasi menjadi SENT_TO_CLIENT.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Final Status: SENT_TO_CLIENT',
        'remarks': 'Memenuhi BR-012, PR-021, PR-022, UAT-1'
    }
]

for idx, test_item in enumerate(system_tests):
    write_test_row(ws_sys, 9 + idx, test_item)

print(f'Populated {len(system_tests)} system test cases.')


# -------------------------------------------------------------
# 3. END-TO-END TESTING - EXECUTION DATA (E2E-001 to E2E-010)
# -------------------------------------------------------------
print('Generating End-to-End Testing test cases...')
ws_e2e = wb['End-to-End Testing - Execution ']
reset_sheet_data(ws_e2e, start_row=9)

e2e_tests = [
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-1: Happy Path Lifecycle',
        'test_id': 'E2E-001',
        'scenario_name': 'Alur Normal Lengkap: Unggah PDF -> AI Ekstrak -> Verifikasi Lolos -> Reviewer Setujui -> Unduh Voucher (SENT_TO_CLIENT)',
        'description': 'Memastikan seluruh siklus 6 status dokumen berjalan mulus dari pengunggahan invoice hingga penerbitan dan pengunduhan voucher resmi.',
        'precondition': 'PIC Klien dan Finance Reviewer aktif, sistem lokal beroperasi penuh',
        'test_steps': (
            '1. PIC Klien login dan mengunggah invoice digital valid (1 hal, Rp 1.500.000, rekening lengkap BCA).\n'
            '2. Sistem menerbitkan Request ID unik dan menetapkan status awal QUEUED.\n'
            '3. BullMQ worker mengambil dokumen dan AI mengekstrak data -> status PROCESSING.\n'
            '4. Sistem menjalankan 3 verifikasi dan menyatakan dokumen bersih -> status READY_FOR_REVIEW.\n'
            '5. Finance Reviewer login, membuka layar split-screen 50:50, memeriksa data, lalu mengklik Setujui.\n'
            '6. Sistem menerbitkan voucher resmi format VCH/YYYYMM/CLI/XXXXX -> status AMIGA_REVIEWED.\n'
            '7. PIC Klien mengunduh voucher PDF resmi -> status dokumen bermutasi menjadi SENT_TO_CLIENT.'
        ),
        'expected_result': 'Seluruh status (QUEUED -> PROCESSING -> READY_FOR_REVIEW -> AMIGA_REVIEWED -> SENT_TO_CLIENT) tercapai tanpa kendala finansial.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'File: invoice_clean_sample.pdf\nClient: CLI-MAIN',
        'remarks': 'Memenuhi UAT-1, BR-001, BR-012'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-2: Multi-line Items Extraction',
        'test_id': 'E2E-002',
        'scenario_name': 'Invoice dengan Multi-line Items (>= 3 Baris) Terbaca Lengkap dan Kalkulasi Total Presisi',
        'description': 'Memastikan dokumen invoice kompleks dengan banyak jenis barang terbaca 100% tanpa baris terlewat dan subtotal terhitung akurat.',
        'precondition': 'Dokumen invoice dengan 3 item barang berbeda',
        'test_steps': (
            '1. Unggah invoice memuat 3 baris item (Item A: 2x 100rb, Item B: 1x 500rb, Item C: 5x 60rb; Subtotal 1jt, PPN 11% 110rb, Total 1.110.000).\n'
            '2. Amati proses ekstraksi AI dan verifikasi matematika sistem.\n'
            '3. Reviewer membuka layar split-screen dan mencocokkan baris item pada PDF dengan tabel editor finansial.'
        ),
        'expected_result': 'Ketiga item terekstraksi presisi, grand total Rp 1.110.000 terverifikasi, status READY_FOR_REVIEW.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Total 3 items, Grand Total: Rp 1.110.000',
        'remarks': 'Memenuhi UAT-2, BR-004'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-3: Missing Evidence Exception',
        'test_id': 'E2E-003',
        'scenario_name': 'Penanganan Invoice Tanpa Rekening Bank (Deteksi MISSING_EVIDENCE & Status EXCEPTION)',
        'description': 'Memastikan sistem mendeteksi ketiadaan rekening transfer, membatalkan review normal, dan menandai dokumen sebagai EXCEPTION.',
        'precondition': 'Berkas invoice sengaja tidak mencantumkan nomor rekening bank',
        'test_steps': (
            '1. PIC Klien mengunggah invoice tanpa nomor rekening transfer.\n'
            '2. AI mengekstrak data teks PDF dan mendapati rekening kosong.\n'
            '3. Verifikasi sistem mengevaluasi ketiadaan rekening.\n'
            '4. Periksa dashboard PIC Klien dan Finance Reviewer.'
        ),
        'expected_result': 'Dokumen masuk status EXCEPTION dengan kode MISSING_EVIDENCE dan tombol koreksi aktif pada sisi PIC Klien.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Rekening: Kosong',
        'remarks': 'Memenuhi UAT-3, BR-008'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-4: Duplicate Detection Exception',
        'test_id': 'E2E-004',
        'scenario_name': 'Pencegahan Pembayaran Ganda (Deteksi POSSIBLE_DUPLICATE & Status EXCEPTION)',
        'description': 'Memastikan sistem mendeteksi upaya pengajuan invoice ganda berdasarkan normalisasi nomor invoice dan nama vendor pada klien yang sama.',
        'precondition': 'Invoice INV-DUPLICATE-001 dari Vendor PT Global telah disetujui sebelumnya',
        'test_steps': (
            '1. Unggah invoice baru dengan nomor "INV-DUPLICATE-001" dan vendor "PT Global" untuk klien CLI-MAIN.\n'
            '2. Tunggu proses ekstraksi dan verifikasi sistem selesai.\n'
            '3. Amati status dokumen dan rincian issue pada dashboard.'
        ),
        'expected_result': 'Sistem mendeteksi kecocokan string ternormalisasi, menandai POSSIBLE_DUPLICATE, dan menetapkan status EXCEPTION.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Duplicate Invoice: INV-DUPLICATE-001',
        'remarks': 'Memenuhi UAT-4, BR-008, SR-013'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-5: Segregation of Duties (SoD)',
        'test_id': 'E2E-005',
        'scenario_name': 'Penegakan SoD Mutlak & Auto-Reassignment ke Reviewer dengan Beban Terendah',
        'description': 'Memastikan jika Reviewer A mengunggah invoice lalu mencoba menyetujui, sistem menolak persetujuan dan mengalihkan penugasan ke Reviewer B.',
        'precondition': 'Reviewer 1 dan Reviewer 2 terdaftar aktif; Reviewer 2 memiliki beban antrean lebih rendah',
        'test_steps': (
            '1. Reviewer 1 mengunggah dokumen invoice.\n'
            '2. Setelah dokumen berstatus READY_FOR_REVIEW, Reviewer 1 membuka URL review dan menekan Setujui.\n'
            '3. Server Fastify mengeksekusi validasi SoD di dalam database transaction atomik serializable.\n'
            '4. Amati respon sistem, status dokumen, dan penugasan baru.'
        ),
        'expected_result': 'Persetujuan ditolak (SOD_VIOLATION 403), dokumen otomatis dialihkan ke Reviewer 2, status tetap READY_FOR_REVIEW, audit log mencatat event.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Uploader: reviewer1@amiga.co.id\nTarget Reassign: reviewer2@amiga.co.id',
        'remarks': 'Memenuhi UAT-5, BR-002, SR-004'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-6: Multi-Tenant Data Isolation',
        'test_id': 'E2E-006',
        'scenario_name': 'Isolasi Ketat Multi-Klien (Penolakan Akses Lintas Tenant / IDOR)',
        'description': 'Memastikan pengguna Klien A dilarang membuka atau memanipulasi dokumen Klien B baik melalui UI maupun panggilan REST API langsung.',
        'precondition': 'Dokumen REQ-CLI2-001 dimiliki oleh Klien CLI-002',
        'test_steps': (
            '1. Login sebagai PIC Klien CLI-MAIN (Klien 1).\n'
            '2. Coba akses endpoint GET /api/v1/invoices/REQ-CLI2-001 milik Klien CLI-002.\n'
            '3. Coba akses voucher download dokumen tersebut.'
        ),
        'expected_result': 'Sistem menolak dengan respon RFC 7807 403 Forbidden atau 404 Not Found, data Klien B tidak bocor.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Cross-client target: REQ-CLI2-001',
        'remarks': 'Memenuhi UAT-6, BR-014, SR-003'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-7: Financial Edit & Audit Trail Diff',
        'test_id': 'E2E-007',
        'scenario_name': 'Penelusuran Audit Immutability: Koreksi Nilai Finansial Tercatat pada Diff Viewer',
        'description': 'Memastikan modifikasi data finansial oleh Reviewer dicatat secara permanen di audit_logs beserta nilai lama, nilai baru, aktor, dan timestamp.',
        'precondition': 'Dokumen berstatus READY_FOR_REVIEW dibuka di layar split-screen',
        'test_steps': (
            '1. Finance Reviewer mengoreksi harga satuan item dari Rp 50.000 menjadi Rp 55.000.\n'
            '2. Sistem menyimpan perubahan dengan penanda is_manually_edited: true.\n'
            '3. Buka halaman audit trail dokumen (/audit/:id).\n'
            '4. Amati visualisasi diff panel perbandingan nilai lama (merah) vs nilai baru (hijau).'
        ),
        'expected_result': 'Diff viewer menampilkan detail perubahan secara kronologis, PII rekening tersensor, dan catatan audit bersifat append-only.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Old: 50.000, New: 55.000',
        'remarks': 'Memenuhi UAT-7, BR-015, PR-025'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'UAT-8: Exception Correction Cycle',
        'test_id': 'E2E-008',
        'scenario_name': 'Perbaikan Dokumen EXCEPTION oleh PIC Klien dengan Mempertahankan Request ID yang Sama',
        'description': 'Memastikan dokumen bermasalah diperbaiki tanpa membuat Request ID baru, diproses ulang, dan berhasil diterbitkan vouchernya.',
        'precondition': 'Dokumen berstatus EXCEPTION dengan Request ID REQ-20261014-00088',
        'test_steps': (
            '1. PIC Klien login dan membuka dokumen berstatus EXCEPTION.\n'
            '2. Buka modal perbaikan dan lengkapi bukti nomor rekening serta nama bank.\n'
            '3. Kirim perbaikan (PUT /api/v1/invoices/:id/correct).\n'
            '4. Verifikasi Request ID pada status PROCESSING dan verifikasi ulang.\n'
            '5. Finance Reviewer menyetujui dokumen yang telah bersih.'
        ),
        'expected_result': 'Request ID TETAP REQ-20261014-00088, dokumen bertransisi ke READY_FOR_REVIEW lalu AMIGA_REVIEWED.',
        'execution_status': 'Not Executed',
        'priority': 'Critical',
        'test_data': 'Request ID identik terverifikasi',
        'remarks': 'Memenuhi UAT-8, BR-009, PR-009'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'Service Lead Workload Rebalance',
        'test_id': 'E2E-009',
        'scenario_name': 'Alur Monitoring & Penyeimbangan Beban Antrean oleh Service Lead',
        'description': 'Memastikan Service Lead dapat mendeteksi penumpukan beban reviewer (>10) dan memindahkan antrean dokumen secara manual.',
        'precondition': 'Reviewer 1 memiliki beban antrean tinggi',
        'test_steps': (
            '1. Service Lead login dan memeriksa bar beban kerja reviewer.\n'
            '2. Memilih dokumen yang menumpuk di antrean Reviewer 1.\n'
            '3. Mengeksekusi aksi Alihkan Penugasan ke Reviewer 2.\n'
            '4. Reviewer 2 login dan memverifikasi dokumen tersebut masuk ke antrean kerjanya.'
        ),
        'expected_result': 'Beban kerja reviewer terdistribusi seimbang, log penugasan tersimpan, dan reviewer 2 dapat memproses dokumen.',
        'execution_status': 'Not Executed',
        'priority': 'Medium',
        'test_data': 'Service Lead Rebalance action',
        'remarks': 'Memenuhi BR-013, PR-023'
    },
    {
        'main_feature': 'End-to-End Flow',
        'sub_feature': 'Reviewer Return Document to PIC',
        'test_id': 'E2E-010',
        'scenario_name': 'Alur Pengembalian Dokumen ke PIC oleh Reviewer dengan Catatan Permasalahan',
        'description': 'Memastikan reviewer dapat mengembalikan invoice yang meragukan/buram ke PIC dengan catatan wajib, dan PIC menerima notifikasi perbaikan.',
        'precondition': 'Dokumen berstatus READY_FOR_REVIEW',
        'test_steps': (
            '1. Finance Reviewer memeriksa dokumen di layar split-screen.\n'
            '2. Menemukan cap invoice menutupi nomor rekening.\n'
            '3. Mengklik Kembalikan ke PIC dan mengisi alasan pengembalian.\n'
            '4. PIC Klien login dan memeriksa dokumen yang berstatus EXCEPTION.'
        ),
        'expected_result': 'Status dokumen berubah menjadi EXCEPTION, catatan reviewer tampil di layar PIC, dan alur perbaikan siap dieksekusi.',
        'execution_status': 'Not Executed',
        'priority': 'High',
        'test_data': 'Alasan: Cap invoice menutupi nomor rekening',
        'remarks': 'Memenuhi PR-017, PR-013'
    }
]

for idx, test_item in enumerate(e2e_tests):
    write_test_row(ws_e2e, 9 + idx, test_item)

print(f'Populated {len(e2e_tests)} E2E test cases.')


# -------------------------------------------------------------
# 4. DEFECT TRACKING DATA (READY FOR LOGGING)
# -------------------------------------------------------------
print('Configuring Defect Tracking Data sheet...')
ws_def = wb['Defect Tracking Data']
reset_sheet_data(ws_def, start_row=9)

# Provide sample defect rows as reference, plus clean rows ready for tester input
sample_defects = [
    {
        'defect_id': 'DEF001',
        'date_detected': datetime.datetime(2026, 10, 14, 0, 0),
        'main_feature': 'Authentication & RBAC',
        'description': 'Pesan validasi kurang ramah pengguna ketika format email salah pada login',
        'precondition': 'Halaman login terbuka',
        'step_by_step': '1. Buka halaman Login\n2. Masukkan email tanpa domain (contoh: user@)\n3. Klik Masuk',
        'defect_status': 'Open',
        'defect_severity': 'Minor',
        'test_data': 'Email: user@',
        'owner': 'QA Engineer',
        'screenshot_link': '-',
        'remarks': 'Perlu menambahkan pesan petunjuk format email standar RFC 5322.'
    },
    {
        'defect_id': 'DEF002',
        'date_detected': datetime.datetime(2026, 10, 14, 0, 0),
        'main_feature': 'Split-Screen Review',
        'description': 'Indikator loading canvas PDF sempat berkedip saat pertama kali berpindah halaman',
        'precondition': 'Layar split-screen terbuka pada invoice 3 halaman',
        'step_by_step': '1. Buka review invoice\n2. Klik navigasi Halaman Berikutnya\n3. Amati canvas render',
        'defect_status': 'Fixed',
        'defect_severity': 'Minor',
        'test_data': 'Multi-page invoice',
        'owner': 'Tech Lead',
        'screenshot_link': '-',
        'remarks': 'Telah diperbaiki dengan debouncing render canvas PDF.js.'
    }
]

def write_defect_row(ws, row_idx, defect_item):
    r = row_idx
    pairs = [
        ('A', 'B'), ('C', 'D'), ('E', 'F'), ('G', 'H'),
        ('I', 'J'), ('K', 'L'), ('M', 'N'), ('O', 'P'),
        ('Q', 'R'), ('S', 'T'), ('U', 'V'), ('W', 'X')
    ]
    for c1, c2 in pairs:
        ws.merge_cells(f'{c1}{r}:{c2}{r}')
        
    ws[f'A{r}'].value = defect_item.get('defect_id', '')
    ws[f'C{r}'].value = defect_item.get('date_detected', datetime.datetime(2026, 10, 14, 0, 0))
    ws[f'E{r}'].value = defect_item.get('main_feature', '')
    ws[f'G{r}'].value = defect_item.get('description', '')
    ws[f'I{r}'].value = defect_item.get('precondition', '')
    ws[f'K{r}'].value = defect_item.get('step_by_step', '')
    ws[f'M{r}'].value = defect_item.get('defect_status', 'New')
    ws[f'O{r}'].value = defect_item.get('defect_severity', 'Minor')
    ws[f'Q{r}'].value = defect_item.get('test_data', '-')
    ws[f'S{r}'].value = defect_item.get('owner', 'QA Engineer')
    ws[f'U{r}'].value = defect_item.get('screenshot_link', '-')
    ws[f'W{r}'].value = defect_item.get('remarks', '-')
    
    center_cols = {'A', 'B', 'C', 'D', 'E', 'F', 'M', 'N', 'O', 'P'}
    for col_idx in range(1, 25):
        c = ws.cell(r, col_idx)
        col_letter = get_column_letter(col_idx)
        c.font = font_regular
        c.border = thin_border
        if col_letter in center_cols:
            c.alignment = align_center
        else:
            c.alignment = align_left
            
    # Defect status fill
    status = defect_item.get('defect_status', '')
    if status == 'Fixed':
        ws.cell(r, 13).fill = PatternFill(start_color='D9EAD3', fill_type='solid')
        ws.cell(r, 14).fill = PatternFill(start_color='D9EAD3', fill_type='solid')
    elif status == 'Open':
        ws.cell(r, 13).fill = PatternFill(start_color='FCE5CD', fill_type='solid')
        ws.cell(r, 14).fill = PatternFill(start_color='FCE5CD', fill_type='solid')
        
    ws.row_dimensions[r].height = 40

for idx, defect in enumerate(sample_defects):
    write_defect_row(ws_def, 9 + idx, defect)

print(f'Configured {len(sample_defects)} initial defect log entries.')

# Save to output path
print(f'Saving final workbook to {OUTPUT_PATH}...')
wb.save(OUTPUT_PATH)
print('Excel Manual Testing Form successfully generated!')
